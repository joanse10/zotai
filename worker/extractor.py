"""
Worker de extragere a textului din documente PDF pentru Zotai.

Strategie în două treceri:
1. Cale rapidă: stratul de text încorporat prin PyMuPDF.
2. Dacă numărul de caractere este sub scan_char_threshold -> documentul este
   considerat scan -> paginile sunt randate în imagini și transmise la
   pytesseract (pachete ron/rus/eng).

Independent de GUI: invocat de orchestrator sau prin `python -m worker`.
"""

import io
import os
import shutil
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field, asdict
from typing import List, Optional

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf


# Căi candidate pentru executabilul Tesseract în ordinea priorității
TESSERACT_CANDIDATES = [
    r"bin\tesseract\tesseract.exe",                       # build portabil în distribuție
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",      # instalare Windows de sistem
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
]

# Diacriticele caracteristice limbii române pentru euristica de detectare a limbii
_RO_DIACRITICS = set("ăâîșțĂÂÎȘȚ")
_RO_STOPWORDS = {"și", "în", "este", "sunt", "pentru", "această", "acest",
                 "unei", "care", "din", "cu", "mai", "sau", "pe", "la"}


@dataclass
class ExtractConfig:
    """Parametrii pipeline-ului de extragere."""
    pages_to_analyze: int = 3        # numărul de pagini inițiale de analizat
    scan_char_threshold: int = 120   # pragul de caractere pentru identificarea scanurilor
    ocr_languages: str = "ron+rus+eng"
    ocr_dpi: int = 300               # DPI de randare pentru OCR
    ocr_all_pages: bool = False      # True — OCR pentru tot documentul, False — primele N pagini
    ocr_workers: int = 4             # threaduri OCR paralele
    tesseract_cmd: str = ""          # cale explicită către tesseract.exe
    tessdata_dir: str = ""           # directorul cu fișiere .traineddata
    max_text_chars: int = 50000      # limita maximă de caractere pentru eșantionul de text


@dataclass
class PageText:
    page: int            # numărul paginii (începând de la 1)
    source: str          # "layer" | "ocr" | "none"
    char_count: int
    text: str


@dataclass
class ExtractResult:
    file_path: str
    file_name: str
    page_count: int = 0
    layer_type: str = ""             # "text" | "scan" | "empty"
    language: str = "UND"            # "RON" | "RUS" | "ENG" | "UND"
    ocr_engine: str = ""             # "tesseract" dacă s-a efectuat OCR
    text: str = ""                   # eșantion de text combinat pentru promptul AI
    pages: List[PageText] = field(default_factory=list)
    ok: bool = True
    error: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


def resolve_tesseract(explicit_cmd: str = "") -> Optional[str]:
    """
    Localizează executabilul Tesseract:
    cale explicită -> variabila de mediu TESSERACT_CMD -> instalare portabilă/sistem -> PATH.
    """
    if explicit_cmd:
        return explicit_cmd if os.path.exists(explicit_cmd) else None

    env_cmd = os.environ.get("TESSERACT_CMD", "")
    if env_cmd and os.path.exists(env_cmd):
        return env_cmd

    app_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    for candidate in TESSERACT_CANDIDATES:
        full = candidate if os.path.isabs(candidate) else os.path.join(app_root, candidate)
        if os.path.exists(full):
            return full

    return shutil.which("tesseract")


def detect_language(text: str) -> str:
    """Detectare euristică a limbii: chirilic -> RUS, diacritice/cuvinte românești -> RON, altfel ENG."""
    letters = sum(1 for c in text if c.isalpha())
    if letters == 0:
        return "UND"

    cyrillic = sum(1 for c in text if "Ѐ" <= c <= "ӿ")
    if cyrillic / letters > 0.15:
        return "RUS"

    diacritics = sum(1 for c in text if c in _RO_DIACRITICS)
    if diacritics / letters > 0.003:
        return "RON"

    words = {w.strip(".,;:!?()[]«»\"'").lower() for w in text.split()}
    if len(words & _RO_STOPWORDS) >= 3:
        return "RON"

    return "ENG"


def _render_page(page, dpi: int):
    """Randează pagina într-o imagine PIL în tonuri de gri pentru OCR."""
    from PIL import Image
    zoom = dpi / 72.0
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csGRAY)
    return Image.open(io.BytesIO(pix.tobytes("png")))


def _ocr_page(pdf_path: str, page_index: int, dpi: int, lang: str) -> PageText:
    """OCR pe o singură pagină. Fiecare thread deschide propriul handle de document."""
    import pytesseract
    with pymupdf.open(pdf_path) as doc:
        img = _render_page(doc[page_index], dpi)
    text = pytesseract.image_to_string(img, lang=lang).strip()
    return PageText(page=page_index + 1, source="ocr", char_count=len(text), text=text)


def _extract_docx(path: str, result: ExtractResult) -> ExtractResult:
    """Extrage textul din paragrafele și tabelele DOCX fără OCR."""
    try:
        import docx
    except ImportError:
        result.ok = False
        result.error = "Extragerea DOCX necesită python-docx (pip install python-docx)"
        return result
    try:
        doc = docx.Document(path)
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.extend(c.text for c in row.cells if c.text.strip())
        text = "\n".join(chunks)
        result.layer_type = "text"
        result.pages = [PageText(page=1, source="layer", char_count=len(text), text=text)]
        result.language = detect_language(text)
        result.text = text[:50000]
        result.page_count = 0
        if not text.strip():
            result.layer_type = "empty"
            result.ok = False
            result.error = "Fișierul DOCX nu conține text"
        return result
    except Exception as e:
        result.ok = False
        result.error = f"Nu s-a putut citi DOCX: {e}"
        return result


def extract_document(pdf_path: str, config: Optional[ExtractConfig] = None) -> ExtractResult:
    """
    Punctul de intrare principal al extracției: identifică tipul stratului și extrage textul.
    PDF digital -> extrage direct stratul de text.
    Scan -> randează paginile și execută Tesseract OCR în paralel.
    DOCX -> extrage textul prin python-docx.
    """
    cfg = config or ExtractConfig()
    result = ExtractResult(
        file_path=pdf_path.replace("\\", "/"),
        file_name=os.path.basename(pdf_path),
    )

    if not os.path.isfile(pdf_path):
        result.ok = False
        result.error = f"Fișierul nu a fost găsit: {pdf_path}"
        return result

    if pdf_path.lower().endswith(".docx"):
        return _extract_docx(pdf_path, result)

    try:
        doc = pymupdf.open(pdf_path)
    except Exception as e:
        result.ok = False
        result.error = f"Nu s-a putut deschide PDF: {e}"
        return result

    with doc:
        result.page_count = doc.page_count
        if result.page_count == 0:
            result.layer_type = "empty"
            result.ok = False
            result.error = "PDF-ul nu conține pagini"
            return result

        # Trecerea 1: stratul de text al primelor N pagini
        analyze_count = min(cfg.pages_to_analyze, result.page_count)
        layer_pages: List[PageText] = []
        for i in range(analyze_count):
            text = doc[i].get_text().strip()
            layer_pages.append(PageText(page=i + 1, source="layer", char_count=len(text), text=text))

        layer_chars = sum(p.char_count for p in layer_pages)
        result.layer_type = "text" if layer_chars >= cfg.scan_char_threshold else "scan"

        if result.layer_type == "text":
            result.pages = layer_pages
        else:
            # Trecerea 2: OCR
            tesseract = resolve_tesseract(cfg.tesseract_cmd)
            if not tesseract:
                result.ok = False
                result.error = (
                    "Documentul este un scan, dar Tesseract nu a fost găsit. "
                    "Specificați calea către tesseract.exe în Setări sau prin --tesseract / TESSERACT_CMD."
                )
                result.pages = layer_pages
                return result

            import pytesseract
            pytesseract.pytesseract.tesseract_cmd = tesseract
            result.ocr_engine = "tesseract"

            tessdata_dir = cfg.tessdata_dir
            if not tessdata_dir:
                from core.paths import tessdata_dir as _tess_dir
                if os.path.isdir(_tess_dir()):
                    tessdata_dir = _tess_dir()
            if tessdata_dir:
                os.environ["TESSDATA_PREFIX"] = tessdata_dir

            ocr_count = result.page_count if cfg.ocr_all_pages else analyze_count
            indices = list(range(ocr_count))
            workers = max(1, min(cfg.ocr_workers, len(indices)))

            from worker.lang_manager import resolve_ocr_languages, installed_languages
            usable_langs = resolve_ocr_languages(cfg.ocr_languages, tessdata_dir)
            if not usable_langs:
                inst = [c for c in installed_languages(tessdata_dir) if c != "osd"]
                if inst:
                    usable_langs = "+".join(inst)
                else:
                    result.ok = False
                    result.error = (
                        "Niciun pachet de limbă OCR nu este instalat în tessdata. "
                        "Descărcați pachete de limbă din Setări -> Limbi OCR."
                    )
                    result.pages = layer_pages
                    return result

            try:
                if workers == 1:
                    result.pages = [_ocr_page(pdf_path, i, cfg.ocr_dpi, usable_langs) for i in indices]
                else:
                    with ThreadPoolExecutor(max_workers=workers) as pool:
                        futures = [pool.submit(_ocr_page, pdf_path, i, cfg.ocr_dpi, usable_langs)
                                   for i in indices]
                        result.pages = [f.result() for f in futures]
            except Exception as e:
                result.ok = False
                result.error = f"Eroare OCR: {e}"
                return result


    combined = "\n\n".join(p.text for p in result.pages if p.text)
    result.text = combined[:cfg.max_text_chars]
    result.language = detect_language(combined)
    return result
